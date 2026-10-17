from datetime import date
from pathlib import Path

import pytest

from wpdc_ingest.core.source import Document, LayoutChanged, Period, RawDocument
from wpdc_ingest.sources.nj_dge_sports_wagering import NJDGESportsWagering, REVENUE_DEFINITION
from wpdc_ingest.sources.uscourts_bankruptcy import USCourtsBankruptcyF2

FIX = Path(__file__).parent / "fixtures"


def raw_for(path: Path, period: Period, kind: str) -> RawDocument:
    return RawDocument(doc=Document(url="https://example.invalid/" + path.name, period=period, kind=kind, label=path.name), path=path, sha256="0" * 64, retrieved_at=date(2026, 10, 16), archive_url=None)


def test_uscourts_f2_parse_and_sum_by_state():
    src = USCourtsBankruptcyF2()
    raw = raw_for(FIX / "uscourts_f2_synthetic.xlsx", Period.quarter_ending(2026, 6), "xlsx")
    recs = src.parse(raw)
    codes = {r["district_code"] for r in recs}
    assert codes == {"OKN", "OKE", "OKW", "KS", "NJ", "ARE", "ND"}  # totals and circuit rows skipped; "OK,N" normalized; "AR \u00b2" footnote stripped and mapped
    nd = next(r for r in recs if r["district_code"] == "ND")
    assert nd["all_nonbusiness"] == 254  # the "2nd" circuit row must not be read as ND
    rows = [r.finalize() for r in src.normalize(recs, raw)]
    ok_all = next(r for r in rows if r["fips"] == "40" and r["chapter"] == "all_nonbusiness")
    assert ok_all["filings"] == 118 + 79 + 196
    ok_13 = next(r for r in rows if r["fips"] == "40" and r["chapter"] == "13")
    assert ok_13["filings"] == 48 + 29 + 75
    assert ok_all["status"] == "pending_review" and ok_all["source_id"] == "uscourts-bankruptcy-f2"
    assert ok_all["notes"] and "F-2" in ok_all["notes"]  # the table reference leads the CSV notes
    assert ok_all["period_start"] == "2026-04-01" and ok_all["period_end"] == "2026-06-30"
    ar = next(r for r in rows if r["fips"] == "05" and r["chapter"] == "all_nonbusiness")
    assert ar["filings"] == 98
    # Twelve-month table: the period is the year ending on the quarter end.
    raw12 = RawDocument(doc=Document(url="u", period=Period.quarter_ending(2026, 6), kind="xlsx", label="12m", meta=(("measure", "twelve_months_ending"),)), path=FIX / "uscourts_f2_synthetic.xlsx", sha256="0" * 64, retrieved_at=date(2026, 10, 16), archive_url=None)
    rows12 = [r.finalize() for r in src.normalize(src.parse(raw12), raw12)]
    assert rows12[0]["period_start"] == "2025-07-01" and rows12[0]["period_end"] == "2026-06-30" and rows12[0]["measure"] == "twelve_months_ending"


def test_uscourts_f2_layout_change_raises(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["Something else entirely"]); ws.append(["District", "Count"])
    p = tmp_path / "bad.xlsx"; wb.save(p)
    with pytest.raises(LayoutChanged):
        USCourtsBankruptcyF2().parse(raw_for(p, Period.quarter_ending(2026, 6), "xlsx"))


def test_nj_dge_parse_totals_and_crosscheck():
    src = NJDGESportsWagering()
    raw = raw_for(FIX / "nj_dge_synthetic.pdf", Period.month(2026, 9), "pdf")
    recs = src.parse(raw)
    by = {r["channel"]: r for r in recs}
    assert by["retail"]["handle_usd"] == 15_000_000 and by["retail"]["gross_revenue_usd"] == 800_000
    assert by["online"]["handle_usd"] == 150_000_000 and by["online"]["gross_revenue_usd"] == 12_500_000
    assert by["all"]["handle_usd"] == 165_000_000
    rows = [r.finalize() for r in src.normalize(recs, raw)]
    assert all(r["revenue_definition"] == REVENUE_DEFINITION for r in rows)
    assert all(r["market_segment"] == "regulated_sportsbook" for r in rows)
    assert {r["channel"] for r in rows} == {"retail", "online", "all"}


def test_nj_dge_wrong_document_raises(tmp_path):
    from tests.make_fixtures import nj_pdf  # reuse writer, then corrupt the text
    p = tmp_path / "x.pdf"; nj_pdf(p)
    data = p.read_bytes().replace(b"SPORTS WAGERING", b"SLOT MACHINES"); p.write_bytes(data)
    with pytest.raises(LayoutChanged):
        NJDGESportsWagering().parse(raw_for(p, Period.month(2026, 9), "pdf"))


def test_writer_upsert_supersedes(tmp_path, monkeypatch):
    from wpdc_ingest.core import writers
    path = tmp_path / "bankruptcy.csv"
    base = {"fips": "40", "period_start": "2026-04-01", "period_end": "2026-06-30", "measure": "quarter", "chapter": "7", "source_id": "uscourts-bankruptcy-f2", "source_url": "u", "archive_url": "", "retrieved_at": "2026-10-16", "status": "pending_review", "entered_by": "t", "reviewed_by": "", "notes": ""}
    key = ("fips", "period_start", "period_end", "measure", "chapter"); sort = ("fips", "period_start", "chapter")
    assert writers.upsert_csv(path, "bankruptcy", [{**base, "filings": 100}], key, sort) == (1, 0)
    assert writers.upsert_csv(path, "bankruptcy", [{**base, "filings": 100}], key, sort) == (0, 0)  # identical: skipped
    assert writers.upsert_csv(path, "bankruptcy", [{**base, "filings": 101}], key, sort) == (1, 1)  # restated: old superseded
    rows = writers.read_csv(path)
    assert sorted(r["status"] for r in rows) == ["pending_review", "superseded"]


def test_uscourts_resolve_finds_f2_tables():
    html = '''<table><tr><td>F-2 Three Months Ended June 30, 2026</td><td><a href="/sites/default/files/data_tables/bf_f2.3_0630.2026.xlsx">Download</a></td></tr>
    <tr><td>F-2 Twelve Months</td><td><a href="https://www.uscourts.gov/sites/default/files/data_tables/bf_f2_0630.2026.xlsx">Download</a></td></tr>
    <tr><td>F-2A Three Months</td><td><a href="/x/bf_f2a_0630.2026.xlsx">Download</a></td></tr>
    <tr><td>F-5A</td><td><a href="/x/bf_f5a_0630.2026.xlsx">Download</a></td></tr></table>'''
    src = USCourtsBankruptcyF2()
    discovered = src.discover(Period.quarter_ending(2026, 6))
    assert discovered[0].url.endswith("/document/bf_f2.3_0630.2026.xlsx") and discovered[0].get("measure") == "quarter"
    assert discovered[1].url.endswith("/document/bf_f2_0630.2026.xlsx") and discovered[1].get("measure") == "twelve_months_ending"
    page = discovered[2]
    assert page.kind == "html" and page.url.endswith("/june-2026-quarterly-bankruptcy-filings")
    docs = src.resolve(page, html)
    assert [d.get("measure") for d in docs] == ["quarter", "twelve_months_ending"]
    assert docs[0].url.startswith("https://www.uscourts.gov/") and docs[0].kind == "xlsx"
    # The real quarter page links to per-table pages rather than files; those map to the file convention.
    docs2 = src.resolve(page, '<a href="/data-news/data-tables/2026/06/30/bankruptcy-filings/f-2">F-2</a><a href="/data-news/data-tables/2026/06/30/bankruptcy-filings/f-2-three-months">F-2 Three Months</a>')
    assert {d.url.rsplit("/", 1)[1] for d in docs2} == {"bf_f2.3_0630.2026.xlsx", "bf_f2_0630.2026.xlsx"}


def test_uscourts_resolve_raises_without_f2():
    with pytest.raises(LayoutChanged):
        USCourtsBankruptcyF2().resolve(Document(url="u", period=Period.quarter_ending(2026, 6), kind="html"), "<a href='/a.pdf'>F-5A</a>")


def test_nj_discover_candidates_and_resolve():
    src = NJDGESportsWagering()
    docs = src.discover(Period.month(2026, 8))
    assert [d.kind for d in docs] == ["html", "html", "html", "pdf"]  # landing candidates, then the historic PDF pattern
    assert docs[-1].url.endswith("SWRTaxReturns/2026/August2026.pdf")
    page = docs[0]
    found = src.resolve(page, '<a href="https://www.nj.gov/oag/ge/docs/Financials/SWRTaxReturns/2026/August2026.pdf">August</a><a href="https://www.nj.gov/oag/ge/docs/Financials/SWRTaxReturns/2026/107AmendmentsList2026.pdf">x</a>')
    assert found[0].url == "https://www.nj.gov/oag/ge/docs/Financials/SWRTaxReturns/2026/August2026.pdf"
    with pytest.raises(LayoutChanged):
        src.resolve(page, "<p>nothing</p>")


def test_census_csv_parse(tmp_path):
    from wpdc_ingest.sources.census_population import CensusPopulation
    csv_text = "SUMLEV,REGION,DIVISION,STATE,NAME,ESTIMATESBASE2020,POPESTIMATE042020,POPESTIMATE2020,POPESTIMATE2021\n010,0,0,00,United States,1,1,2,3\n040,3,6,40,Oklahoma,3959353,3959353,3964912,3986639\n040,X,X,72,Puerto Rico,1,1,2,3\n"
    p = tmp_path / "nst.csv"; p.write_text(csv_text)
    src = CensusPopulation()
    doc = src.discover(None)[0]
    raw = RawDocument(doc=Document(url=doc.url, kind="csv", label=doc.label, meta=(("vintage", 2025),)), path=p, sha256="0" * 64, retrieved_at=date(2026, 10, 16), archive_url=None)
    recs = src.parse(raw)
    assert [(r["fips"], r["year"], r["population"]) for r in recs] == [("40", 2020, 3964912), ("40", 2021, 3986639)]
    assert all(r["vintage"] == 2025 for r in recs)


def test_census_api_html_error_raises(tmp_path):
    from wpdc_ingest.sources.census_population import CensusPopulation
    p = tmp_path / "x.json"; p.write_text("<html>missing key</html>")
    src = CensusPopulation()
    doc = [d for d in src.discover(None) if d.kind == "json"][0]
    with pytest.raises(LayoutChanged):
        src.parse(RawDocument(doc=doc, path=p, sha256="0" * 64, retrieved_at=date(2026, 10, 16), archive_url=None))


def test_nj_tax_returns_sum_across_licensees_and_skip_skin_pages():
    src = NJDGESportsWagering()
    raw = raw_for(FIX / "nj_dge107_synthetic.pdf", Period.month(2026, 8), "pdf")
    recs = {r["channel"]: r for r in src.parse(raw)}
    assert recs["retail"]["gross_revenue_usd"] == 3_668 + 100_000
    assert recs["retail"]["promo_deductions_usd"] == -2_000  # parentheses are negative, dash is zero
    assert recs["retail"]["state_tax_usd"] == 312 + 8_330
    assert recs["online"]["gross_revenue_usd"] == 5_000_000 + 3_000_000  # skin detail page not double counted
    assert recs["online"]["state_tax_usd"] == 987_500 + 592_500
    assert recs["all"]["gross_revenue_usd"] == 3_668 + 100_000 + 8_000_000
    assert recs["retail"]["handle_usd"] is None
    rows = [r.finalize() for r in src.normalize(list(recs.values()), raw)]
    assert all(r["revenue_definition"].startswith("Gross revenue before deductions") for r in rows)


def test_nj_tax_returns_wrong_month_raises(tmp_path):
    src = NJDGESportsWagering()
    with pytest.raises(LayoutChanged):
        src.parse(raw_for(FIX / "nj_dge107_synthetic.pdf", Period.month(2026, 9), "pdf"))


def test_nj_tax_returns_2018_form_family_a():
    src = NJDGESportsWagering()
    raw = raw_for(FIX / "nj_dge107_2018_synthetic.pdf", Period.month(2018, 9), "pdf")
    recs = {r["channel"]: r for r in src.parse(raw)}
    assert recs["retail"]["gross_revenue_usd"] == 394_046 - 40_848  # line 1, parentheses negative
    assert recs["online"]["gross_revenue_usd"] == 106_463 + 223_339  # line 8
    assert recs["retail"]["state_tax_usd"] == 33_494 + 1_611 and recs["online"]["state_tax_usd"] == 13_840 + 29_035
    assert recs["retail"]["adjusted_revenue_usd"] is None and recs["all"]["promo_deductions_usd"] is None  # not on this form
    assert "SEPTMEBER" in recs["all"]["notes"]  # misspelt month tolerated and recorded
    rows = [r.finalize() for r in src.normalize(list(recs.values()), raw)]
    assert "2018 to May 2019 form" in rows[0]["notes"] and "SEPTMEBER" in rows[0]["notes"]


def test_nj_tax_returns_2019_form_family_b_split_digits():
    src = NJDGESportsWagering()
    raw = raw_for(FIX / "nj_dge107_2019_synthetic.pdf", Period.month(2019, 12), "pdf")
    recs = {r["channel"]: r for r in src.parse(raw)}
    assert recs["retail"]["gross_revenue_usd"] == 108_437 and recs["retail"]["adjusted_revenue_usd"] == 108_437
    assert recs["online"]["gross_revenue_usd"] == 26_752  # "2 6,752" read as one number; footed against year to date
    assert recs["online"]["state_tax_usd"] == 3_478
    assert recs["all"]["gross_revenue_usd"] == 108_437 + 26_752


def test_nj_tax_returns_misfooting_form_raises():
    with pytest.raises(LayoutChanged):
        NJDGESportsWagering().parse(raw_for(FIX / "nj_dge107_misfoot_synthetic.pdf", Period.month(2019, 12), "pdf"))


def test_uscourts_f2_pdf_2022_layout():
    src = USCourtsBankruptcyF2()
    recs = {r["district_code"]: r for r in src.parse(raw_for(FIX / "uscourts_f2_2022_synthetic.pdf", Period.quarter_ending(2022, 12), "pdf"))}
    assert set(recs) == {"NJ", "ARE", "ND", "OKN"}  # Total, circuit and Virgin Islands rows skipped
    assert recs["NJ"] == {"district_code": "NJ", "all_nonbusiness": 2415, "7": 1349, "11": 3, "13": 1063}
    assert recs["ND"]["all_nonbusiness"] == 97 and recs["ARE"]["13"] == 879


def test_uscourts_f2_pdf_2018_layout_repairs_glued_last_column():
    src = USCourtsBankruptcyF2()
    recs = {r["district_code"]: r for r in src.parse(raw_for(FIX / "uscourts_f2_2018_synthetic.pdf", Period.quarter_ending(2018, 3), "pdf"))}
    assert set(recs) == {"NJ", "ARE", "ARW", "ND", "OKN"}
    assert recs["NJ"] == {"district_code": "NJ", "all_nonbusiness": 6485, "7": 3967, "11": 29, "13": 2489}  # "2,4890" repaired
    assert recs["OKN"]["13"] == 67 and recs["ND"]["13"] == 11 and recs["ARW"]["13"] == 403
    rows = [r.finalize() for r in src.normalize(list(recs.values()), raw_for(FIX / "uscourts_f2_2018_synthetic.pdf", Period.quarter_ending(2018, 3), "pdf"))]
    ar = next(r for r in rows if r["fips"] == "05" and r["chapter"] == "13")
    assert ar["filings"] == 949 + 403 and "PDF edition" in ar["notes"]


def test_uscourts_resolve_takes_real_links_and_prefers_xlsx():
    src = USCourtsBankruptcyF2()
    page = Document(url="u", period=Period.quarter_ending(2022, 12), kind="html")
    html = '<a href="/sites/default/files/data_tables/bf_f2.3_1231.2022.pdf">x</a><a href="/sites/default/files/data_tables/bf_f2_1231.2022.pdf">y</a><a href="/sites/default/files/data_tables/bf_f2_1231.2022.xlsx">z</a>'
    docs = {d.get("measure"): d for d in src.resolve(page, html)}
    assert docs["quarter"].url.endswith("bf_f2.3_1231.2022.pdf") and docs["quarter"].kind == "pdf"
    assert docs["twelve_months_ending"].url.endswith("bf_f2_1231.2022.xlsx")


def test_uscourts_f2_pdf_doubled_header_page_is_read():
    src = USCourtsBankruptcyF2()
    recs = {r["district_code"]: r for r in src.parse(raw_for(FIX / "uscourts_f2_doubled_synthetic.pdf", Period.quarter_ending(2021, 12), "pdf"))}
    assert set(recs) == {"DC", "ME"} and recs["ME"]["all_nonbusiness"] == 147 and recs["DC"]["13"] == 11


def test_uscourts_resolve_accepts_reuploaded_file_suffix():
    src = USCourtsBankruptcyF2()
    page = Document(url="u", period=Period.quarter_ending(2018, 6), kind="html")
    html = '<a href="/data-news/data-tables/2018/06/30/bankruptcy-filings/f-2-three-months">x</a><a href="/sites/default/files/bf_f2.3_0630.2018_0.pdf">y</a><a href="/sites/default/files/bf_f2_0630.2018.pdf">z</a>'
    docs = {d.get("measure"): d for d in src.resolve(page, html)}
    assert docs["quarter"].url.endswith("/bf_f2.3_0630.2018_0.pdf") and docs["twelve_months_ending"].url.endswith("/bf_f2_0630.2018.pdf")


def test_archive_docs_extension_from_type_url_or_bytes():
    from wpdc_ingest.archive_docs import extension
    assert extension("https://x/a.xlsx", "application/octet-stream", b"PK\x03\x04xl/workbook.xml") == "xlsx"
    assert extension("https://x/report?id=1", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", b"PK") == "xlsx"
    assert extension("https://x/page?f=1", "text/html; charset=utf-8", b"<!DOCTYPE html>") == "html"
    assert extension("https://x/r.pdf", "", b"%PDF-1.7") == "pdf"
    assert extension("https://x/d", "text/csv", b"a,b") == "csv"
    assert extension("https://x/d", "", b"\x00\x01") == "bin"


def test_ny_statewide_monthly_report_parses_every_month_and_crosschecks_totals():
    from wpdc_ingest.sources.ny_gaming_commission_mobile_sports import NYGamingCommissionMobileSports, REVENUE_DEFINITION
    src = NYGamingCommissionMobileSports()
    raw = raw_for(FIX / "ny_monthly_mobile_statewide_2026-09.pdf", None, "pdf")
    recs = src.parse(raw)
    assert len(recs) == 56  # January 2022 to August 2026
    aug26 = next(r for r in recs if (r["year"], r["month"]) == (2026, 8))
    assert (aug26["handle_usd"], aug26["gross_revenue_usd"], aug26["state_tax_usd"]) == (1_731_302_503, 168_476_334, 85_934_232)
    apr23 = next(r for r in recs if (r["year"], r["month"]) == (2023, 4))  # row with a blank adjustments cell
    assert (apr23["handle_usd"], apr23["gross_revenue_usd"], apr23["state_tax_usd"]) == (1_544_465_378, 138_820_564, 70_798_487)
    jan22 = next(r for r in recs if (r["year"], r["month"]) == (2022, 1))  # first month, on the page with Unclaimed Funds
    assert jan22["handle_usd"] == 1_669_719_034 and jan22["state_tax_usd"] == 63_310_867
    rows = [r.finalize() for r in src.normalize(recs, raw)]
    assert len(rows) == 56 and all(r["channel"] == "online" and r["fips"] == "36" for r in rows)
    first = min(rows, key=lambda r: r["period_start"])
    assert first["period_start"] == "2022-01-01" and first["period_end"] == "2022-01-31"
    assert rows[0]["revenue_definition"] == REVENUE_DEFINITION and rows[0]["adjusted_revenue_usd"] == rows[0]["gross_revenue_usd"]
    assert "page" in rows[0]["notes"]
    # a single requested month filters to that month; a month not in the edition fails loudly
    one = src.normalize(recs, raw_for(FIX / "ny_monthly_mobile_statewide_2026-09.pdf", Period.month(2025, 11), "pdf"))
    assert len(one) == 1 and one[0].values["handle_usd"] == 2_606_597_963
    with pytest.raises(LayoutChanged):
        src.normalize(recs, raw_for(FIX / "ny_monthly_mobile_statewide_2026-09.pdf", Period.month(2026, 9), "pdf"))


def test_ny_report_layout_change_raises(tmp_path):
    from wpdc_ingest.sources.ny_gaming_commission_mobile_sports import NYGamingCommissionMobileSports
    with pytest.raises(LayoutChanged):
        NYGamingCommissionMobileSports().parse(raw_for(FIX / "nj_dge_synthetic.pdf", None, "pdf"))


def test_pa_workbook_grand_total_months_and_crosschecks():
    from wpdc_ingest.sources.pa_pgcb_sports_wagering import PAPGCBSportsWagering, REVENUE_DEFINITION
    src = PAPGCBSportsWagering()
    raw = raw_for(FIX / "pa_pgcb_sports_wagering_fy2024-25.xlsx", None, "xlsx")
    recs = src.parse(raw)
    assert [(r["year"], r["month"]) for r in recs][:2] == [(2024, 7), (2024, 8)] and len(recs) == 12
    jul = recs[0]
    assert round(jul["total"]["handle"]) == 414_673_804 and round(jul["online"]["promo"]) == 14_161_688 and round(jul["retail"]["taxable"]) == 3_244_022
    rows = [r.finalize() for r in src.normalize(recs, raw)]
    assert len(rows) == 36
    a = next(r for r in rows if r["period_start"] == "2024-07-01" and r["channel"] == "all")
    assert (a["handle_usd"], a["gross_revenue_usd"], a["promo_deductions_usd"], a["adjusted_revenue_usd"], a["state_tax_usd"]) == (414_673_804, 52_727_378, 14_161_688, 38_565_690, 13_112_334)
    r = next(r for r in rows if r["period_start"] == "2024-07-01" and r["channel"] == "retail")
    assert r["gross_revenue_usd"] == r["adjusted_revenue_usd"] == 3_244_022 and r["promo_deductions_usd"] == 0
    assert a["revenue_definition"] == REVENUE_DEFINITION and "GRAND TOTAL" in a["notes"]
    # the older workbook naming carries the same layout
    old = src.parse(raw_for(FIX / "pa_pgcb_sports_wagering_fy2022-23.xlsx", None, "xlsx"))
    assert len(old) == 12 and round(old[0]["total"]["handle"]) == 336_507_932
    # a requested month filters, a missing one fails
    one = src.normalize(recs, raw_for(FIX / "pa_pgcb_sports_wagering_fy2024-25.xlsx", Period.month(2025, 1), "xlsx"))
    assert len(one) == 3 and one[0].values["period_start"] == "2025-01-01"
    with pytest.raises(LayoutChanged):
        src.normalize(recs, raw_for(FIX / "pa_pgcb_sports_wagering_fy2024-25.xlsx", Period.month(2023, 1), "xlsx"))


def test_pa_layout_change_raises(tmp_path):
    import openpyxl
    from wpdc_ingest.sources.pa_pgcb_sports_wagering import PAPGCBSportsWagering
    wb = openpyxl.Workbook(); ws = wb.active
    for _ in range(12): ws.append(["something", 1, 2])
    p = tmp_path / "pa.xlsx"; wb.save(p)
    with pytest.raises(LayoutChanged):
        PAPGCBSportsWagering().parse(raw_for(p, None, "xlsx"))


def test_ks_monthly_summary_parses_and_crosschecks():
    from wpdc_ingest.sources.ks_lottery_sports_wagering import KSLotterySportsWagering, REVENUE_DEFINITION
    src = KSLotterySportsWagering()
    raw = raw_for(FIX / "ks_lottery_sports_wagering_2026-01.pdf", Period.month(2026, 1), "pdf")
    recs = {r["channel"]: r for r in src.parse(raw)}
    assert (recs["retail"]["handle_usd"], recs["retail"]["adjusted_revenue_usd"], recs["retail"]["state_tax_usd"]) == (4_200_982, 688_350, 68_835)
    assert (recs["online"]["handle_usd"], recs["online"]["adjusted_revenue_usd"]) == (223_109_404, 14_509_154)
    assert (recs["all"]["handle_usd"], recs["all"]["state_tax_usd"]) == (227_310_386, 1_519_750)
    rows = [r.finalize() for r in src.normalize(list(recs.values()), raw)]
    assert len(rows) == 3 and rows[0]["fips"] == "20" and rows[0]["period_start"] == "2026-01-01" and rows[0]["gross_revenue_usd"] is None
    assert rows[0]["revenue_definition"] == REVENUE_DEFINITION
    # an older edition (BetMGM at zero) parses the same way
    old = {r["channel"]: r for r in src.parse(raw_for(FIX / "ks_lottery_sports_wagering_2023-05.pdf", Period.month(2023, 5), "pdf"))}
    assert old["all"]["handle_usd"] > 0 and old["retail"]["handle_usd"] + old["online"]["handle_usd"] == old["all"]["handle_usd"]
    with pytest.raises(LayoutChanged):  # wrong month requested for this file
        src.parse(raw_for(FIX / "ks_lottery_sports_wagering_2026-01.pdf", Period.month(2026, 2), "pdf"))
    # resolve picks the month's link from the publications page
    html = '<a href="/media/jswjvfff/sports-wagering-monthly-revenue-2026-01.pdf">January</a><a href="/media/04weixyd/sports-wagering-monthly-revenue-2026-02.pdf">February</a>'
    docs = src.resolve(Document(url="u", period=Period.month(2026, 2), kind="html"), html)
    assert docs[0].url.endswith("2026-02.pdf")
    assert len(src.resolve(Document(url="u", period=None, kind="html"), html)) == 2


def test_oh_repair_numbers_rejoins_split_digits():
    from wpdc_ingest.sources.oh_occc_sports_gaming import repair_numbers
    assert repair_numbers("917,625,932 8 15,852,742 3,442,131 30,951,463 98,331,059 9 8,331,059") == [917_625_932, 815_852_742, 3_442_131, 30_951_463, 98_331_059, 98_331_059]
    assert repair_numbers("13,524,481 1 2,177,453 4 9,952 400 1 ,297,076 1,354,989") == [13_524_481, 12_177_453, 49_952, 400, 1_297_076, 1_354_989]
    assert repair_numbers("532,852 5 89,243 1,522 - (57,913) -") == [532_852, 589_243, 1_522, 0, -57_913, 0]


def test_oh_yearly_report_parses_months_and_crosschecks():
    from wpdc_ingest.sources.oh_occc_sports_gaming import OHOCCCSportsGaming, REVENUE_DEFINITION
    src = OHOCCCSportsGaming()
    raw = raw_for(FIX / "oh_occc_sports_gaming_2026_report04.pdf", None, "pdf")
    recs = src.parse(raw)
    assert sorted({r["month"] for r in recs}) == [1, 2, 3, 4] and len(recs) == 12
    jan_all = next(r for r in recs if r["month"] == 1 and r["channel"] == "all")
    assert (jan_all["handle_usd"], jan_all["gross_revenue_usd"], jan_all["adjusted_revenue_usd"]) == (931_150_413, 99_628_135, 99_686_048)
    feb_retail = next(r for r in recs if r["month"] == 2 and r["channel"] == "retail")
    assert (feb_retail["handle_usd"], feb_retail["gross_revenue_usd"]) == (10_261_446, 372_550)
    rows = [r.finalize() for r in src.normalize(recs, raw)]
    assert len(rows) == 12 and rows[0]["fips"] == "39" and rows[0]["state_tax_usd"] is None and rows[0]["revenue_definition"] == REVENUE_DEFINITION
    one = src.normalize(recs, raw_for(FIX / "oh_occc_sports_gaming_2026_report04.pdf", Period.month(2026, 4), "pdf"))
    assert len(one) == 3 and one[0].values["period_start"] == "2026-04-01"
    with pytest.raises(LayoutChanged):
        src.normalize(recs, raw_for(FIX / "oh_occc_sports_gaming_2026_report04.pdf", Period.month(2026, 5), "pdf"))
    with pytest.raises(LayoutChanged):
        src.parse(raw_for(FIX / "ks_lottery_sports_wagering_2026-01.pdf", None, "pdf"))


def test_pa_first_fiscal_year_has_eight_months_and_no_pre_promo_revenue():
    from wpdc_ingest.sources.pa_pgcb_sports_wagering import PAPGCBSportsWagering
    src = PAPGCBSportsWagering()
    raw = raw_for(FIX / "pa_pgcb_sports_wagering_fy2018-19.xlsx", None, "xlsx")
    recs = src.parse(raw)
    assert [(r["year"], r["month"]) for r in recs][0] == (2018, 11) and len(recs) == 8
    rows = [r.finalize() for r in src.normalize(recs, raw)]
    nov_all = next(r for r in rows if r["period_start"] == "2018-11-01" and r["channel"] == "all")
    assert nov_all["handle_usd"] == 1_414_587 and nov_all["adjusted_revenue_usd"] == 508_997 and nov_all["gross_revenue_usd"] is None and nov_all["promo_deductions_usd"] is None
    nov_online = next(r for r in rows if r["period_start"] == "2018-11-01" and r["channel"] == "online")
    assert nov_online["handle_usd"] == 0


def test_mi_workbook_statewide_group_and_operator_crosscheck():
    from wpdc_ingest.sources.mi_mgcb_internet_sports_betting import MIMGCBInternetSportsBetting, REVENUE_DEFINITION
    src = MIMGCBInternetSportsBetting()
    raw = raw_for(FIX / "mi_mgcb_internet_sports_betting_2026.xlsx", None, "xlsx")
    recs = src.parse(raw)
    assert [(r["year"], r["month"]) for r in recs] == [(2026, m) for m in range(1, 9)]
    jan = recs[0]
    assert round(jan["handle_usd"]) == 491_260_369 and round(jan["gross_revenue_usd"]) == 58_036_388 and round(jan["adjusted_revenue_usd"]) == 36_965_495 and round(jan["state_tax_usd"]) == 2_538_064
    rows = [r.finalize() for r in src.normalize(recs, raw)]
    assert len(rows) == 8 and rows[0]["channel"] == "online" and rows[0]["fips"] == "26" and rows[0]["revenue_definition"] == REVENUE_DEFINITION
    first = src.parse(raw_for(FIX / "mi_mgcb_internet_sports_betting_2021.xlsx", None, "xlsx"))
    assert len(first) == 12 and first[0]["month"] == 1
    html = '<a href="/mgcb/-/media/x/Internet-Sports-Betting---2026.xlsx?rev=1&amp;hash=2">2026 Internet Sports Betting Revenues and Taxes&nbsp;Excel</a><a href="/a.pdf">2026 Internet Sports Betting Revenues and Taxes PDF</a>'
    docs = src.resolve(Document(url="u", period=Period.month(2026, 3), kind="html"), html)
    assert len(docs) == 1 and docs[0].url.startswith("https://www.michigan.gov/mgcb/-/media/x/Internet-Sports-Betting---2026.xlsx?rev=1&hash=2")
    with pytest.raises(LayoutChanged):
        src.parse(raw_for(FIX / "pa_pgcb_sports_wagering_fy2024-25.xlsx", None, "xlsx"))


def test_ks_wide_spaced_edition_is_rebuilt_from_character_positions():
    from wpdc_ingest.sources.ks_lottery_sports_wagering import KSLotterySportsWagering
    src = KSLotterySportsWagering()
    recs = {r["channel"]: r for r in src.parse(raw_for(FIX / "ks_lottery_sports_wagering_2025-04.pdf", Period.month(2025, 4), "pdf"))}
    assert recs["all"]["handle_usd"] == recs["retail"]["handle_usd"] + recs["online"]["handle_usd"] > 100_000_000
    assert recs["all"]["state_tax_usd"] > 0


def test_oh_full_year_report_with_negative_cells_parses():
    from wpdc_ingest.sources.oh_occc_sports_gaming import OHOCCCSportsGaming, repair_numbers
    assert repair_numbers("1 3,374,834 13,410,654 86,468 - ( 122,288) 6 13,206") == [13_374_834, 13_410_654, 86_468, 0, -122_288, 613_206]
    recs = OHOCCCSportsGaming().parse(raw_for(FIX / "oh_occc_sports_gaming_2025_report12.pdf", None, "pdf"))
    assert sorted({r["month"] for r in recs}) == list(range(1, 13)) and len(recs) == 36


def test_provenance_block_falls_back_to_the_content_hash():
    from wpdc_ingest.core import provenance
    raw = raw_for(FIX / "ks_lottery_sports_wagering_2026-01.pdf", Period.month(2026, 1), "pdf")
    assert provenance.block("x", raw, method="pdf", entered_by="t")["archive_url"] == "sha256:" + "0" * 64


def test_il_two_part_month_file_parses_and_crosschecks():
    from wpdc_ingest.sources.il_igb_sports_wagering import ILIGBSportsWagering, REVENUE_DEFINITION, _forms
    src = ILIGBSportsWagering()
    raw = raw_for(FIX / "il_igb_sports_wagering_2026-08.txt", Period.month(2026, 8), "txt")
    recs = {r["channel"]: r for r in src.parse(raw)}
    assert recs["all"]["handle_usd"] == recs["retail"]["handle_usd"] + recs["online"]["handle_usd"] > 500_000_000
    assert recs["all"]["agr"] > 50_000_000 and recs["online"]["agr"] > recs["retail"]["agr"]
    assert round(recs["retail"]["agr"]) == 1_123_840  # the PDF edition's in-person AGR total
    rows = [r.finalize() for r in src.normalize(list(recs.values()), raw)]
    assert len(rows) == 3 and rows[0]["fips"] == "17" and rows[0]["gross_revenue_usd"] == rows[0]["adjusted_revenue_usd"] and rows[0]["revenue_definition"] == REVENUE_DEFINITION
    with pytest.raises(LayoutChanged):
        src.parse(raw_for(FIX / "il_igb_sports_wagering_2026-08.txt", Period.month(2026, 7), "txt"))
    docs = src.discover(Period.month(2024, 5))
    assert len(docs) == 1 and docs[0].get("measure") == "2024-05" and "Tax Summary Report" in docs[0].get("posts")
    assert _forms(2024, 5)[1]["data"]["SearchCash"] == "Detail Report" and _forms(2024, 5)[0]["data"]["SearchEndMonth"] == "May"
    series = src.discover(None)
    assert series[0].period.start.isoformat() == "2020-03-01" and len(series) > 70


def test_pa_current_workbook_skips_months_not_yet_reported():
    from wpdc_ingest.sources.pa_pgcb_sports_wagering import PAPGCBSportsWagering
    recs = PAPGCBSportsWagering().parse(raw_for(FIX / "pa_pgcb_sports_wagering_fy2026-27.xlsx", None, "xlsx"))
    assert [(r["year"], r["month"]) for r in recs] == [(2026, 7), (2026, 8)]
